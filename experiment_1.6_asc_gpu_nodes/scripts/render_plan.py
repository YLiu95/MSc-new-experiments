from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt


def inline_text(tokens) -> str:
    return "".join("\n" if token.type in ("softbreak", "hardbreak") else token.content
                   for token in tokens if token.type in ("text", "code_inline", "softbreak", "hardbreak"))


def add_inline(paragraph, tokens) -> None:
    bold = False
    italic = False
    for token in tokens:
        if token.type == "strong_open":
            bold = True
        elif token.type == "strong_close":
            bold = False
        elif token.type == "em_open":
            italic = True
        elif token.type == "em_close":
            italic = False
        elif token.type in ("text", "code_inline", "softbreak", "hardbreak"):
            run = paragraph.add_run("\n" if token.type in ("softbreak", "hardbreak") else token.content)
            run.bold = bold
            run.italic = italic
            if token.type == "code_inline":
                run.font.name = "Consolas"


def render(source: Path, target: Path) -> None:
    tokens = MarkdownIt("commonmark").enable("table").parse(source.read_text())
    document = Document()
    document.sections[0].top_margin = Inches(0.75)
    document.sections[0].bottom_margin = Inches(0.75)
    document.styles["Normal"].font.size = Pt(9)
    document.styles["Normal"].font.name = "Calibri"
    paragraph = None
    heading = None
    list_styles = []
    table_rows = None
    row = None
    cell = None
    for token in tokens:
        if token.type == "heading_open":
            heading = f"Heading {min(9, int(token.tag[1:]))}"
            paragraph = document.add_paragraph(style=heading)
        elif token.type == "heading_close":
            paragraph = None
            heading = None
        elif token.type in ("bullet_list_open", "ordered_list_open"):
            list_styles.append("List Bullet" if token.type == "bullet_list_open" else "List Number")
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            list_styles.pop()
        elif token.type == "paragraph_open" and table_rows is None and heading is None:
            paragraph = document.add_paragraph(style=list_styles[-1] if list_styles else "Normal")
        elif token.type == "paragraph_close" and heading is None:
            paragraph = None
        elif token.type == "table_open":
            table_rows = []
        elif token.type == "tr_open":
            row = []
        elif token.type in ("th_open", "td_open"):
            cell = ""
        elif token.type in ("th_close", "td_close"):
            row.append(cell)
            cell = None
        elif token.type == "tr_close":
            table_rows.append(row)
            row = None
        elif token.type == "table_close":
            if table_rows:
                table = document.add_table(rows=0, cols=len(table_rows[0]))
                table.style = "Table Grid"
                for values in table_rows:
                    if len(values) != len(table_rows[0]):
                        raise ValueError("Markdown table has inconsistent column counts")
                    for column, text in zip(table.add_row().cells, values):
                        column.text = text
            table_rows = None
        elif token.type == "inline":
            if cell is not None:
                cell += inline_text(token.children)
            elif paragraph is not None:
                add_inline(paragraph, token.children)
        elif token.type in ("fence", "code_block"):
            block = document.add_paragraph(style="No Spacing")
            code = block.add_run(token.content)
            code.font.name = "Consolas"
            code.font.size = Pt(8)
        elif token.type == "hr":
            document.add_paragraph("---")
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    arguments = parser.parse_args()
    render(arguments.source, arguments.output)


if __name__ == "__main__":
    main()